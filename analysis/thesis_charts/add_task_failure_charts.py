from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


WORK_DIR = Path("/Users/justintiensmith/Documents/lerobot/output/thesis_bar_charts_audit")
OUTPUT_DIR = Path(
    "/Users/justintiensmith/Documents/lerobot/outputs/"
    "01a0678e-4e7e-7a31-b018-7fba72ab9fe5"
)
SOURCE = OUTPUT_DIR / "MSc_Thesis_Bar_Charts_with_Model_Failure_Profiles.xlsx"
OUTPUT = OUTPUT_DIR / "MSc_Thesis_Bar_Charts_with_Task_Failure_Profiles.xlsx"
DATA = json.loads((WORK_DIR / "task_failure_profile_data.json").read_text(encoding="utf-8"))

DATA_SHEET = "Task Failure Data"
CHART_SHEETS = {
    "Novel Spatial Grounding": "Novel Failures by Task",
    "Familiar Spatial Grounding": "Familiar Failures by Task",
}

HEADER_BLUE = "1F4E78"
TEXT = "1F1F1F"
SUBTLE = "595959"
FAILURE_COLORS = {
    "Wrong target": "4E79A7",
    "No action": "F28E2B",
    "Grasp failure": "59A14F",
    "Correct target, incorrect spatial outcome": "E15759",
    "Timeout or incomplete execution": "EDC948",
    "Premature release or dropped object": "B07AA1",
    "Control drift or instability": "9C755F",
    "Release/retention failure": "76B7B2",
}
LEGEND_LABELS = {
    "Wrong target": "Wrong target",
    "No action": "No action",
    "Grasp failure": "Grasp failure",
    "Correct target, incorrect spatial outcome": "Incorrect spatial outcome",
    "Timeout or incomplete execution": "Timeout/incomplete",
    "Premature release or dropped object": "Premature release/drop",
    "Control drift or instability": "Control drift",
    "Release/retention failure": "Release/retention",
}
MODEL_LABELS = {
    "VLA-0": "VLA-0",
    "GR00T Full Fine-Tune": "GR00T Full FT",
    "GR00T Frozen LLM": "GR00T Frozen LLM",
    "Pi0.5": "Pi0.5",
    "MolmoAct2": "MolmoAct2",
}


def set_row(ws, row: int, start_col: int, values: list[object]) -> None:
    for offset, value in enumerate(values):
        ws.cell(row=row, column=start_col + offset, value=value)


def chart_signature(chart) -> dict:
    anchor = chart.anchor
    anchor_signature = None
    if hasattr(anchor, "_from"):
        anchor_signature = (
            anchor._from.col,
            anchor._from.colOff,
            anchor._from.row,
            anchor._from.rowOff,
        )
    series = []
    for item in chart.ser:
        value_formula = item.val.numRef.f if item.val and item.val.numRef else None
        category_formula = None
        if item.cat:
            if item.cat.strRef:
                category_formula = item.cat.strRef.f
            elif item.cat.numRef:
                category_formula = item.cat.numRef.f
            elif item.cat.multiLvlStrRef:
                category_formula = item.cat.multiLvlStrRef.f
        series.append((value_formula, category_formula))
    return {
        "class": type(chart).__name__,
        "type": getattr(chart, "type", None),
        "grouping": getattr(chart, "grouping", None),
        "width": chart.width,
        "height": chart.height,
        "anchor": anchor_signature,
        "series": series,
    }


def worksheet_signature(ws) -> dict:
    cells = [
        (row, column, cell.value, cell.data_type, cell.style_id, cell.number_format)
        for (row, column), cell in sorted(ws._cells.items())
    ]
    return {
        "cells": cells,
        "charts": [chart_signature(chart) for chart in ws._charts],
        "sheet_format": (ws.sheet_format.defaultColWidth, ws.sheet_format.defaultRowHeight),
        "freeze": str(ws.freeze_panes),
        "gridlines": ws.sheet_view.showGridLines,
    }


def get_rows(condition: str, task: str, model: str) -> list[dict]:
    rows = [
        row
        for row in DATA["rows"]
        if row["Condition"] == condition and row["Task"] == task and row["Model"] == model
    ]
    order = {mode: index for index, mode in enumerate(DATA["failure_mode_order"])}
    rows.sort(key=lambda row: order[row["Failure Mode"]])
    if len(rows) != len(DATA["failure_mode_order"]):
        raise AssertionError(f"Incomplete failure profile: {condition}, {task}, {model}")
    return rows


wb = load_workbook(SOURCE)
original_sheet_names = list(wb.sheetnames)
original_signatures = {name: worksheet_signature(wb[name]) for name in original_sheet_names}

for name in [DATA_SHEET, *CHART_SHEETS.values()]:
    if name in wb.sheetnames:
        del wb[name]

# Flat audit table plus wide chart matrices.
data_ws = wb.create_sheet(DATA_SHEET)
data_ws.sheet_view.showGridLines = False
data_ws["A1"] = "Task-level failure composition data"
data_ws["A2"] = (
    "Shares and Wilson intervals use the failed rollouts in each condition-task-model group as the denominator."
)
data_ws["A3"] = (
    "These values describe the composition of failures, not the probability of failure over all attempts."
)
data_ws["A4"] = "Source: Spa_Bench_All_Failure_Audit.xlsx, All Classified Failures."
data_ws["A1"].font = Font(name="Arial", size=16, bold=True, color=TEXT)
for row in (2, 3, 4):
    data_ws.cell(row, 1).font = Font(name="Arial", size=10, italic=True, color=SUBTLE)

headers = [
    "Condition",
    "Task",
    "Model",
    "Failure mode",
    "Count",
    "Task-model failures",
    "Share",
    "Wilson 95% lower",
    "Wilson 95% upper",
    "Lower error",
    "Upper error",
]
set_row(data_ws, 6, 1, headers)
header_fill = PatternFill("solid", fgColor=HEADER_BLUE)
for cell in data_ws[6][:11]:
    cell.fill = header_fill
    cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
data_ws.row_dimensions[6].height = 30

thin = Side(style="thin", color="D9E2F3")
flat_row = 7
for condition in CHART_SHEETS:
    for task in DATA["task_order"]:
        for model in DATA["model_order"][condition]:
            for datum in get_rows(condition, task, model):
                set_row(
                    data_ws,
                    flat_row,
                    1,
                    [
                        datum["Condition"],
                        datum["Task"],
                        MODEL_LABELS[datum["Model"]],
                        datum["Failure Mode"],
                        datum["Count"],
                        datum["Task-Model Failures"],
                        datum["Share"],
                        datum["CI Lower"],
                        datum["CI Upper"],
                        datum["Lower Error"],
                        datum["Upper Error"],
                    ],
                )
                for cell in data_ws[flat_row][:11]:
                    cell.font = Font(name="Arial", size=10, color=TEXT)
                    cell.border = Border(bottom=thin)
                for column in (5, 6):
                    data_ws.cell(flat_row, column).number_format = "#,##0"
                for column in range(7, 12):
                    data_ws.cell(flat_row, column).number_format = "0.0%"
                flat_row += 1

widths = {
    "A": 27,
    "B": 25,
    "C": 24,
    "D": 44,
    "E": 10,
    "F": 18,
    "G": 12,
    "H": 18,
    "I": 18,
    "J": 14,
    "K": 14,
}
for column, width in widths.items():
    data_ws.column_dimensions[column].width = width
data_ws.freeze_panes = "A7"
data_ws.auto_filter.ref = f"A6:K{flat_row - 1}"

data_ws["M1"] = "Chart matrices"
data_ws["M1"].font = Font(name="Arial", size=13, bold=True, color=TEXT)
data_ws["M2"] = "Counts are normalized to 100% within each model bar by the chart."
data_ws["M2"].font = Font(name="Arial", size=10, italic=True, color=SUBTLE)
data_ws.column_dimensions["M"].width = 29
for column in ("N", "O", "P", "Q", "R", "S", "T", "U"):
    data_ws.column_dimensions[column].width = 24

matrix_records: dict[str, dict[str, dict]] = {condition: {} for condition in CHART_SHEETS}
matrix_row = 5
for condition in CHART_SHEETS:
    for task in DATA["task_order"]:
        title_row = matrix_row
        header_row = matrix_row + 1
        models = DATA["model_order"][condition]
        first_model_row = matrix_row + 2
        last_model_row = first_model_row + len(models) - 1
        data_ws.cell(title_row, 13, f"{condition}: {task}")
        data_ws.cell(title_row, 13).font = Font(name="Arial", size=11, bold=True, color=TEXT)
        set_row(data_ws, header_row, 13, ["Model", *DATA["failure_mode_order"]])
        for cell in data_ws[header_row][12:21]:
            cell.fill = header_fill
            cell.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        data_ws.row_dimensions[header_row].height = 42

        for row_number, model in enumerate(models, start=first_model_row):
            rows = get_rows(condition, task, model)
            total = rows[0]["Task-Model Failures"]
            set_row(
                data_ws,
                row_number,
                13,
                [
                    f"{MODEL_LABELS[model]} (n={total})",
                    *[row["Count"] for row in rows],
                ],
            )
            for cell in data_ws[row_number][12:21]:
                cell.font = Font(name="Arial", size=9, color=TEXT)
                cell.border = Border(bottom=thin)
            for column in range(14, 22):
                data_ws.cell(row_number, column).number_format = "#,##0"

        matrix_records[condition][task] = {
            "header_row": header_row,
            "first_model_row": first_model_row,
            "last_model_row": last_model_row,
            "models": models,
        }
        matrix_row = last_model_row + 3


def add_shared_legend(ws) -> None:
    legend_positions = [
        (5, 1),
        (5, 5),
        (5, 10),
        (5, 14),
        (6, 1),
        (6, 5),
        (6, 10),
        (6, 14),
    ]
    for mode, (row, column) in zip(DATA["failure_mode_order"], legend_positions):
        swatch = ws.cell(row, column)
        swatch.fill = PatternFill("solid", fgColor=FAILURE_COLORS[mode])
        swatch.border = Border(
            left=Side(style="thin", color=FAILURE_COLORS[mode]),
            right=Side(style="thin", color=FAILURE_COLORS[mode]),
            top=Side(style="thin", color=FAILURE_COLORS[mode]),
            bottom=Side(style="thin", color=FAILURE_COLORS[mode]),
        )
        label = ws.cell(row, column + 1, LEGEND_LABELS[mode])
        label.font = Font(name="Arial", size=9, color=TEXT)
        label.alignment = Alignment(vertical="center")
    for column in ("A", "E", "J", "N"):
        ws.column_dimensions[column].width = 3
    for column in ("B", "F", "K", "O"):
        ws.column_dimensions[column].width = 20


chart_positions = ["A8", "J8", "A29", "J29", "A50", "J50"]
for condition, sheet_name in CHART_SHEETS.items():
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False
    ws["A1"] = f"{condition} failure composition by task and model"
    ws["A2"] = (
        "Each 100%-stacked bar shows the distribution of primary failure modes within one task-model group's failed rollouts."
    )
    ws["A3"] = (
        "The n in each model label is the number of failures. Small n produces unstable shares; Wilson intervals are in Task Failure Data."
    )
    ws["A4"] = "Source: Spa_Bench_All_Failure_Audit.xlsx, All Classified Failures."
    ws["A1"].font = Font(name="Arial", size=16, bold=True, color=TEXT)
    for row in (2, 3, 4):
        ws.cell(row, 1).font = Font(name="Arial", size=10, italic=True, color=SUBTLE)
    add_shared_legend(ws)

    for task, position in zip(DATA["task_order"], chart_positions):
        record = matrix_records[condition][task]
        chart = BarChart()
        chart.type = "bar"
        chart.grouping = "percentStacked"
        chart.overlap = 100
        chart.gapWidth = 45 if condition == "Familiar Spatial Grounding" else 70
        chart.style = 10
        chart.title = task
        chart.legend = None
        chart.width = 14.2
        chart.height = 7.6
        chart.x_axis.scaling.orientation = "maxMin"
        chart.x_axis.axPos = "l"
        chart.y_axis.title = "Share of failures"
        chart.y_axis.scaling.min = 0
        chart.y_axis.scaling.max = 1
        chart.y_axis.majorUnit = 0.5
        chart.y_axis.numFmt = "0%"
        chart.y_axis.axPos = "b"
        chart.y_axis.tickLblPos = "low"
        values = Reference(
            data_ws,
            min_col=14,
            max_col=21,
            min_row=record["header_row"],
            max_row=record["last_model_row"],
        )
        categories = Reference(
            data_ws,
            min_col=13,
            min_row=record["first_model_row"],
            max_row=record["last_model_row"],
        )
        chart.add_data(values, titles_from_data=True, from_rows=False)
        chart.set_categories(categories)
        if len(chart.ser) != len(DATA["failure_mode_order"]):
            raise AssertionError(f"Unexpected series count for {condition}, {task}")
        for series, mode in zip(chart.ser, DATA["failure_mode_order"]):
            series.graphicalProperties.solidFill = FAILURE_COLORS[mode]
            series.graphicalProperties.line.solidFill = FAILURE_COLORS[mode]
        ws.add_chart(chart, position)

    ws.print_area = "A1:R69"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.35

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"
wb.save(OUTPUT)


# Verify the saved workbook and protect every pre-existing sheet.
verified = load_workbook(OUTPUT, data_only=False)
for name in original_sheet_names:
    if worksheet_signature(verified[name]) != original_signatures[name]:
        raise AssertionError(f"Existing sheet changed unexpectedly: {name}")

saved_data_ws = verified[DATA_SHEET]
flat_count_total = sum(
    saved_data_ws.cell(row, 5).value or 0 for row in range(7, flat_row)
)
expected_total = sum(DATA["condition_totals"].values())
if flat_count_total != expected_total:
    raise AssertionError(
        f"Task-level failure counts do not reconcile: expected {expected_total}, found {flat_count_total}"
    )

chart_checks = {}
for condition, sheet_name in CHART_SHEETS.items():
    saved_ws = verified[sheet_name]
    if len(saved_ws._charts) != 6:
        raise AssertionError(f"Expected six charts in {sheet_name}; found {len(saved_ws._charts)}")
    chart_checks[sheet_name] = []
    for task, chart in zip(DATA["task_order"], saved_ws._charts):
        if chart.type != "bar" or chart.grouping != "percentStacked":
            raise AssertionError(f"Unexpected chart type for {sheet_name}, {task}")
        if chart.y_axis.scaling.min != 0 or chart.y_axis.scaling.max != 1:
            raise AssertionError(f"Unexpected axis scale for {sheet_name}, {task}")
        if len(chart.ser) != len(DATA["failure_mode_order"]):
            raise AssertionError(f"Unexpected series count for {sheet_name}, {task}")
        colors = [series.graphicalProperties.solidFill.srgbClr for series in chart.ser]
        expected_colors = [FAILURE_COLORS[mode] for mode in DATA["failure_mode_order"]]
        if colors != expected_colors:
            raise AssertionError(f"Failure-mode colors are inconsistent for {sheet_name}, {task}")
        chart_checks[sheet_name].append(
            {
                "task": task,
                "series": len(chart.ser),
                "categories": len(DATA["model_order"][condition]),
                "scale": [0, 1],
            }
        )

formula_error_tokens = {
    "#REF!",
    "#DIV/0!",
    "#VALUE!",
    "#NAME?",
    "#N/A",
    "#NUM!",
    "#NULL!",
    "#SPILL!",
    "#CALC!",
}
formula_errors = []
for ws in verified.worksheets:
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and any(token in cell.value for token in formula_error_tokens):
                formula_errors.append((ws.title, cell.coordinate, cell.value))
if formula_errors:
    raise AssertionError(f"Formula-error text found: {formula_errors[:10]}")

verification = {
    "source": str(SOURCE),
    "output": str(OUTPUT),
    "new_sheets": [*CHART_SHEETS.values(), DATA_SHEET],
    "condition_totals": DATA["condition_totals"],
    "task_model_totals": DATA["task_model_totals"],
    "flat_rows": flat_row - 7,
    "flat_count_total": flat_count_total,
    "chart_checks": chart_checks,
    "existing_sheets_unchanged": True,
    "formula_errors": [],
}
(WORK_DIR / "task_failure_chart_verification.json").write_text(
    json.dumps(verification, indent=2), encoding="utf-8"
)
print(json.dumps(verification, indent=2))
