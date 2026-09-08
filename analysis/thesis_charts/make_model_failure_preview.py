from copy import copy, deepcopy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.pagebreak import Break


SOURCE = Path(
    "/Users/justintiensmith/Documents/lerobot/outputs/"
    "01a0678e-4e7e-7a31-b018-7fba72ab9fe5/"
    "MSc_Thesis_Bar_Charts_with_Model_Failure_Profiles.xlsx"
)
OUTPUT = Path("/private/tmp/spa_bench_model_failure_preview.xlsx")
CHART_SHEET = "Failure by Model - All"
DATA_SHEET = "Failure by Model - Data"

source_book = load_workbook(SOURCE)
book = Workbook()
chart_sheet = book.active
chart_sheet.title = CHART_SHEET
data_sheet = book.create_sheet(DATA_SHEET)


def copy_cells_and_dimensions(source_sheet, target_sheet) -> None:
    target_sheet.sheet_view.showGridLines = source_sheet.sheet_view.showGridLines
    target_sheet.freeze_panes = source_sheet.freeze_panes
    for (row, column), source_cell in source_sheet._cells.items():
        target_cell = target_sheet.cell(row=row, column=column, value=source_cell.value)
        if source_cell.has_style:
            target_cell._style = copy(source_cell._style)
        if source_cell.number_format:
            target_cell.number_format = source_cell.number_format
        if source_cell.alignment:
            target_cell.alignment = copy(source_cell.alignment)
        if source_cell.font:
            target_cell.font = copy(source_cell.font)
        if source_cell.fill:
            target_cell.fill = copy(source_cell.fill)
        if source_cell.border:
            target_cell.border = copy(source_cell.border)
    for key, dimension in source_sheet.column_dimensions.items():
        target_sheet.column_dimensions[key].width = dimension.width
    for key, dimension in source_sheet.row_dimensions.items():
        target_sheet.row_dimensions[key].height = dimension.height


copy_cells_and_dimensions(source_book[CHART_SHEET], chart_sheet)
copy_cells_and_dimensions(source_book[DATA_SHEET], data_sheet)
for chart in source_book[CHART_SHEET]._charts:
    chart_sheet.add_chart(deepcopy(chart), chart.anchor)
data_sheet.sheet_state = "hidden"

chart_sheet.print_area = "A1:N143"
chart_sheet.page_setup.orientation = "landscape"
chart_sheet.page_setup.paperSize = chart_sheet.PAPERSIZE_A4
chart_sheet.page_setup.fitToWidth = 1
chart_sheet.page_setup.fitToHeight = 0
chart_sheet.sheet_properties.pageSetUpPr.fitToPage = True
chart_sheet.page_margins.left = 0.25
chart_sheet.page_margins.right = 0.25
chart_sheet.page_margins.top = 0.35
chart_sheet.page_margins.bottom = 0.35
for row_id in (33, 61, 89, 117):
    chart_sheet.row_breaks.append(Break(id=row_id))

book.save(OUTPUT)
print(OUTPUT)
