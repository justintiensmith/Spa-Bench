from copy import copy, deepcopy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.pagebreak import Break


SOURCE = Path(
    "/Users/justintiensmith/Documents/lerobot/outputs/"
    "01a0678e-4e7e-7a31-b018-7fba72ab9fe5/"
    "MSc_Thesis_Bar_Charts_with_Task_Failure_Profiles.xlsx"
)
OUTPUT = Path("/private/tmp/spa_bench_task_failure_preview.xlsx")
CHART_SHEETS = ["Novel Failures by Task", "Familiar Failures by Task"]
DATA_SHEET = "Task Failure Data"


def copy_cells_and_dimensions(source_sheet, target_sheet) -> None:
    target_sheet.sheet_view.showGridLines = source_sheet.sheet_view.showGridLines
    target_sheet.freeze_panes = source_sheet.freeze_panes
    for (row, column), source_cell in source_sheet._cells.items():
        target_cell = target_sheet.cell(row=row, column=column, value=source_cell.value)
        if source_cell.has_style:
            target_cell._style = copy(source_cell._style)
        if source_cell.number_format:
            target_cell.number_format = source_cell.number_format
        if source_cell.alignment:
            target_cell.alignment = copy(source_cell.alignment)
        if source_cell.font:
            target_cell.font = copy(source_cell.font)
        if source_cell.fill:
            target_cell.fill = copy(source_cell.fill)
        if source_cell.border:
            target_cell.border = copy(source_cell.border)
    for key, dimension in source_sheet.column_dimensions.items():
        target_sheet.column_dimensions[key].width = dimension.width
    for key, dimension in source_sheet.row_dimensions.items():
        target_sheet.row_dimensions[key].height = dimension.height


source_book = load_workbook(SOURCE)
preview_book = Workbook()
preview_book.remove(preview_book.active)

for sheet_name in CHART_SHEETS:
    source_sheet = source_book[sheet_name]
    target_sheet = preview_book.create_sheet(sheet_name)
    copy_cells_and_dimensions(source_sheet, target_sheet)
    for chart in source_sheet._charts:
        target_sheet.add_chart(deepcopy(chart), chart.anchor)
    target_sheet.print_area = "A1:R69"
    target_sheet.page_setup.orientation = "landscape"
    target_sheet.page_setup.paperSize = target_sheet.PAPERSIZE_A3
    target_sheet.page_setup.fitToWidth = 1
    target_sheet.page_setup.fitToHeight = 0
    target_sheet.sheet_properties.pageSetUpPr.fitToPage = True
    for row_id in (28, 49):
        target_sheet.row_breaks.append(Break(id=row_id))

data_sheet = preview_book.create_sheet(DATA_SHEET)
copy_cells_and_dimensions(source_book[DATA_SHEET], data_sheet)
data_sheet.sheet_state = "hidden"

preview_book.save(OUTPUT)
print(OUTPUT)
