from copy import copy, deepcopy
from pathlib import Path

from openpyxl import Workbook, load_workbook


SOURCE = Path(
    "/Users/justintiensmith/Documents/lerobot/outputs/"
    "01a0678e-4e7e-7a31-b018-7fba72ab9fe5/"
    "MSc_Thesis_Bar_Charts_with_Novel_Failure_Counts.xlsx"
)
OUTPUT = Path("/private/tmp/spa_bench_novel_failure_count_preview.xlsx")
SHEET_NAME = "Novel Failure Counts by Task"

source_book = load_workbook(SOURCE)
source_sheet = source_book[SHEET_NAME]
book = Workbook()
sheet = book.active
sheet.title = SHEET_NAME
sheet.sheet_view.showGridLines = source_sheet.sheet_view.showGridLines

for (row, column), source_cell in source_sheet._cells.items():
    target_cell = sheet.cell(row=row, column=column, value=source_cell.value)
    if source_cell.has_style:
        target_cell._style = copy(source_cell._style)
    if source_cell.number_format:
        target_cell.number_format = source_cell.number_format
    if source_cell.alignment:
        target_cell.alignment = copy(source_cell.alignment)
    if source_cell.font:
        target_cell.font = copy(source_cell.font)
    if source_cell.fill:
        target_cell.fill = copy(source_cell.fill)
    if source_cell.border:
        target_cell.border = copy(source_cell.border)

for key, dimension in source_sheet.column_dimensions.items():
    sheet.column_dimensions[key].width = dimension.width
for key, dimension in source_sheet.row_dimensions.items():
    sheet.row_dimensions[key].height = dimension.height
for chart in source_sheet._charts:
    sheet.add_chart(deepcopy(chart), chart.anchor)

sheet.print_area = "A1:R44"
sheet.page_setup.orientation = "landscape"
sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
sheet.page_setup.fitToWidth = 1
sheet.page_setup.fitToHeight = 1
sheet.sheet_properties.pageSetUpPr.fitToPage = True
sheet.page_margins.left = 0.25
sheet.page_margins.right = 0.25
sheet.page_margins.top = 0.35
sheet.page_margins.bottom = 0.35

book.save(OUTPUT)
print(OUTPUT)
