from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


WORK_DIR = Path("/Users/justintiensmith/Documents/lerobot/output/thesis_bar_charts_audit")
OUTPUT_DIR = Path(
    "/Users/justintiensmith/Documents/lerobot/outputs/"
    "01a0678e-4e7e-7a31-b018-7fba72ab9fe5"
)
SOURCE = OUTPUT_DIR / "MSc_Thesis_Bar_Charts_with_Task_Failure_Profiles.xlsx"
OUTPUT = OUTPUT_DIR / "MSc_Thesis_Bar_Charts_with_Novel_Failure_Counts.xlsx"
DATA = json.loads((WORK_DIR / "task_failure_profile_data.json").read_text(encoding="utf-8"))

SHEET_NAME = "Novel Failure Counts by Task"
CONDITION = "Novel Spatial Grounding"
MODELS = ["GR00T Frozen LLM", "Pi0.5", "MolmoAct2"]
MODEL_COLORS = {
    "GR00T Frozen LLM": "D95140",
    "Pi0.5": "59A55B",
    "MolmoAct2": "5383ED",
}
HEADER_BLUE = "1F4E78"
TEXT = "1F1F1F"
SUBTLE = "595959"


def set_row(ws, row: int, values: list[object]) -> None:
    for column, value in enumerate(values, start=1):
        ws.cell(row=row, column=column, value=value)


def chart_signature(chart) -> dict:
    anchor = chart.anchor
    anchor_signature = None
    if hasattr(anchor, "_from"):
        anchor_signature = (
            anchor._from.col,
            anchor._from.colOff,
            anchor._from.row,
            anchor._from.rowOff,
        )
    series = []
    for item in chart.ser:
        value_formula = item.val.numRef.f if item.val and item.val.numRef else None
        category_formula = None
        if item.cat:
            if item.cat.strRef:
                category_formula = item.cat.strRef.f
            elif item.cat.numRef:
                category_formula = item.cat.numRef.f
            elif item.cat.multiLvlStrRef:
                category_formula = item.cat.multiLvlStrRef.f
        series.append((value_formula, category_formula))
    return {
        "class": type(chart).__name__,
        "type": getattr(chart, "type", None),
        "grouping": getattr(chart, "grouping", None),
        "width": chart.width,
        "height": chart.height,
        "anchor": anchor_signature,
        "series": series,
    }


def worksheet_signature(ws) -> dict:
    cells = [
        (row, column, cell.value, cell.data_type, cell.style_id, cell.number_format)
        for (row, column), cell in sorted(ws._cells.items())
    ]
    return {
        "cells": cells,
        "charts": [chart_signature(chart) for chart in ws._charts],
        "sheet_format": (ws.sheet_format.defaultColWidth, ws.sheet_format.defaultRowHeight),
        "freeze": str(ws.freeze_panes),
        "gridlines": ws.sheet_view.showGridLines,
    }


wb = load_workbook(SOURCE)
original_sheet_names = list(wb.sheetnames)
original_signatures = {name: worksheet_signature(wb[name]) for name in original_sheet_names}

if SHEET_NAME in wb.sheetnames:
    del wb[SHEET_NAME]
ws = wb.create_sheet(SHEET_NAME)
ws.sheet_view.showGridLines = False

ws["A1"] = "Novel spatial grounding failures by task"
ws["A2"] = (
    "Each task contains 50 trials per model. Bar height is the total number of failed rollouts across the three models."
)
ws["A3"] = "Colored segments show each model's contribution; the fixed maximum is 150 failed rollouts per task."
ws["A4"] = "Source: Spa_Bench_All_Failure_Audit.xlsx, All Classified Failures."
ws["A1"].font = Font(name="Arial", size=16, bold=True, color=TEXT)
for row in (2, 3, 4):
    ws.cell(row, 1).font = Font(name="Arial", size=10, italic=True, color=SUBTLE)

headers = [
    "Task",
    "GR00T Frozen LLM",
    "Pi0.5",
    "MolmoAct2",
    "Total failures",
    "Total trials",
    "Failure rate",
]
set_row(ws, 6, headers)
header_fill = PatternFill("solid", fgColor=HEADER_BLUE)
for cell in ws[6][:7]:
    cell.fill = header_fill
    cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws.row_dimensions[6].height = 30

thin = Side(style="thin", color="D9E2F3")
task_totals = []
for row_number, task in enumerate(DATA["task_order"], start=7):
    counts = [DATA["task_model_totals"][CONDITION][task][model] for model in MODELS]
    task_total = sum(counts)
    task_totals.append(task_total)
    set_row(ws, row_number, [task, *counts, task_total, 150, task_total / 150])
    for cell in ws[row_number][:7]:
        cell.font = Font(name="Arial", size=10, color=TEXT)
        cell.border = Border(bottom=thin)
    for column in range(2, 7):
        ws.cell(row_number, column).number_format = "#,##0"
    ws.cell(row_number, 7).number_format = "0.0%"

model_totals = [
    sum(DATA["task_model_totals"][CONDITION][task][model] for task in DATA["task_order"])
    for model in MODELS
]
grand_total = sum(model_totals)
set_row(ws, 13, ["Across six tasks", *model_totals, grand_total, 900, grand_total / 900])
for cell in ws[13][:7]:
    cell.font = Font(name="Arial", size=10, bold=True, color=TEXT)
    cell.border = Border(top=Side(style="medium", color=HEADER_BLUE))
for column in range(2, 7):
    ws.cell(13, column).number_format = "#,##0"
ws.cell(13, 7).number_format = "0.0%"

ws.column_dimensions["A"].width = 28
for column in ("B", "C", "D"):
    ws.column_dimensions[column].width = 21
for column in ("E", "F", "G"):
    ws.column_dimensions[column].width = 16

chart = BarChart()
chart.type = "col"
chart.grouping = "stacked"
chart.overlap = 100
chart.gapWidth = 55
chart.style = 10
chart.title = "Novel spatial grounding failures by task"
chart.width = 25
chart.height = 13
chart.y_axis.title = "Failed rollouts"
chart.y_axis.scaling.min = 0
chart.y_axis.scaling.max = 150
chart.y_axis.majorUnit = 25
chart.y_axis.numFmt = "0"
chart.x_axis.title = "Task"
chart.legend.position = "b"

values = Reference(ws, min_col=2, max_col=4, min_row=6, max_row=12)
categories = Reference(ws, min_col=1, min_row=7, max_row=12)
chart.add_data(values, titles_from_data=True, from_rows=False)
chart.set_categories(categories)
for series, model in zip(chart.ser, MODELS):
    series.graphicalProperties.solidFill = MODEL_COLORS[model]
    series.graphicalProperties.line.solidFill = MODEL_COLORS[model]
chart.dLbls = DataLabelList(
    showVal=True,
    showCatName=False,
    showSerName=False,
    showLegendKey=False,
    showPercent=False,
    showBubbleSize=False,
    showLeaderLines=False,
    numFmt="0",
    dLblPos="ctr",
)
ws.add_chart(chart, "A16")

ws.print_area = "A1:R44"
ws.page_setup.orientation = "landscape"
ws.page_setup.paperSize = ws.PAPERSIZE_A3
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_margins.left = 0.25
ws.page_margins.right = 0.25
ws.page_margins.top = 0.35
ws.page_margins.bottom = 0.35

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"
wb.save(OUTPUT)


# Verify the saved workbook and protect every pre-existing sheet.
verified = load_workbook(OUTPUT, data_only=False)
for name in original_sheet_names:
    if worksheet_signature(verified[name]) != original_signatures[name]:
        raise AssertionError(f"Existing sheet changed unexpectedly: {name}")

saved_ws = verified[SHEET_NAME]
if len(saved_ws._charts) != 1:
    raise AssertionError(f"Expected one chart; found {len(saved_ws._charts)}")
saved_chart = saved_ws._charts[0]
if saved_chart.type != "col" or saved_chart.grouping != "stacked":
    raise AssertionError("Expected a stacked column chart.")
if len(saved_chart.ser) != 3:
    raise AssertionError(f"Expected three model series; found {len(saved_chart.ser)}")
if saved_chart.y_axis.scaling.min != 0 or saved_chart.y_axis.scaling.max != 150:
    raise AssertionError("Expected a common 0-150 failure-count scale.")

saved_colors = [series.graphicalProperties.solidFill.srgbClr for series in saved_chart.ser]
expected_colors = [MODEL_COLORS[model] for model in MODELS]
if saved_colors != expected_colors:
    raise AssertionError(f"Model colors do not match the existing workbook: {saved_colors}")

saved_counts = [
    [saved_ws.cell(row, column).value for column in range(2, 5)]
    for row in range(7, 13)
]
expected_counts = [
    [DATA["task_model_totals"][CONDITION][task][model] for model in MODELS]
    for task in DATA["task_order"]
]
if saved_counts != expected_counts:
    raise AssertionError("Saved failure counts do not match the verified task totals.")
if [sum(row) for row in saved_counts] != task_totals:
    raise AssertionError("Saved task totals do not reconcile.")
if [sum(row[index] for row in saved_counts) for index in range(3)] != model_totals:
    raise AssertionError("Saved model totals do not reconcile.")
if sum(task_totals) != 489 or grand_total != 489:
    raise AssertionError("Grand total must reconcile to 489 novel-condition failures.")

formula_error_tokens = {
    "#REF!",
    "#DIV/0!",
    "#VALUE!",
    "#NAME?",
    "#N/A",
    "#NUM!",
    "#NULL!",
    "#SPILL!",
    "#CALC!",
}
formula_errors = []
for worksheet in verified.worksheets:
    for row in worksheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and any(token in cell.value for token in formula_error_tokens):
                formula_errors.append((worksheet.title, cell.coordinate, cell.value))
if formula_errors:
    raise AssertionError(f"Formula-error text found: {formula_errors[:10]}")

verification = {
    "source": str(SOURCE),
    "output": str(OUTPUT),
    "new_sheet": SHEET_NAME,
    "task_order": DATA["task_order"],
    "failure_counts": expected_counts,
    "task_totals": task_totals,
    "model_totals": dict(zip(MODELS, model_totals)),
    "grand_total": grand_total,
    "chart_type": "stacked column",
    "chart_scale": [0, 150],
    "model_colors": MODEL_COLORS,
    "existing_sheets_unchanged": True,
    "formula_errors": [],
}
(WORK_DIR / "novel_failure_count_chart_verification.json").write_text(
    json.dumps(verification, indent=2), encoding="utf-8"
)
print(json.dumps(verification, indent=2))
