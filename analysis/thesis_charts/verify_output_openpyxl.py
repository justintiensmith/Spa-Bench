from __future__ import annotations

import json
import zipfile
from pathlib import Path

from openpyxl import load_workbook


OUTPUT = Path(
    "/Users/justintiensmith/Documents/lerobot/outputs/"
    "01a0678e-4e7e-7a31-b018-7fba72ab9fe5/"
    "MSc_Thesis_Bar_Charts_Verified_with_CIs.xlsx"
)
WORK_DIR = Path(__file__).resolve().parent


with zipfile.ZipFile(OUTPUT) as archive:
    bad_zip_member = archive.testzip()

wb = load_workbook(OUTPUT, data_only=False)
target_sheets = [
    "Paraphrasing Comparison v3",
    "No Movement",
    "Relational Movement No Movement",
    "Cross Task Object Chart v2",
    "Familiar Spatial Task Chart",
    "Matched Manipulation Task Chart",
    "Novel Spatial Task Chart",
    "Emergent Skills Data",
    "Sheet2",
    "Failure Analysis - Novel",
    "Failure Analysis - All",
]

chart_checks = {}
for sheet_name in target_sheets:
    ws = wb[sheet_name]
    sheet_charts = []
    for chart in ws._charts:
        series_checks = []
        for series in chart.ser:
            bars = series.errBars
            series_checks.append(
                {
                    "value_formula": series.val.numRef.f if series.val and series.val.numRef else None,
                    "error_type": bars.errValType if bars else None,
                    "error_direction": bars.errDir if bars else None,
                    "minus_formula": bars.minus.numRef.f if bars and bars.minus and bars.minus.numRef else None,
                    "plus_formula": bars.plus.numRef.f if bars and bars.plus and bars.plus.numRef else None,
                }
            )
        sheet_charts.append(series_checks)
    chart_checks[sheet_name] = sheet_charts

all_target_series_have_custom_ci = all(
    series["error_type"] == "cust" and series["minus_formula"] and series["plus_formula"]
    for charts in chart_checks.values()
    for series_group in charts
    for series in series_group
)

formula_errors = []
for ws in wb.worksheets:
    for cell in ws._cells.values():
        if isinstance(cell.value, str) and cell.value.startswith(("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!")):
            formula_errors.append(f"{ws.title}!{cell.coordinate}: {cell.value}")

report = {
    "zip_integrity_error": bad_zip_member,
    "sheets": wb.sheetnames,
    "all_target_series_have_custom_ci": all_target_series_have_custom_ci,
    "formula_errors": formula_errors,
    "chart_checks": chart_checks,
    "failure_chart_blue": {
        sheet: wb[sheet]._charts[0].ser[0].graphicalProperties.solidFill
        for sheet in ["Failure Analysis - Novel", "Failure Analysis - All"]
    },
}
(WORK_DIR / "openpyxl_verification.json").write_text(json.dumps(report, indent=2, default=str), encoding="utf-8")
print(json.dumps({key: value for key, value in report.items() if key != "chart_checks"}, indent=2, default=str))
if bad_zip_member or formula_errors or not all_target_series_have_custom_ci:
    raise SystemExit(1)
