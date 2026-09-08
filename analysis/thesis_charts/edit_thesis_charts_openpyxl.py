from __future__ import annotations

import json
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.data_source import NumData, NumDataSource, NumRef, NumVal, StrData, StrVal
from openpyxl.chart.error_bar import ErrorBars
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.cell import range_boundaries


SOURCE = Path("/Users/justintiensmith/Documents/MSc_Thesis_Bar_Charts.xlsx")
WORK_DIR = Path("/Users/justintiensmith/Documents/lerobot/output/thesis_bar_charts_audit")
OUTPUT_DIR = Path(
    "/Users/justintiensmith/Documents/lerobot/outputs/"
    "01a0678e-4e7e-7a31-b018-7fba72ab9fe5"
)
OUTPUT = OUTPUT_DIR / "MSc_Thesis_Bar_Charts_Verified_with_CIs.xlsx"

EXPECTED = json.loads((WORK_DIR / "expected_chart_data.json").read_text(encoding="utf-8"))
FAILURES = json.loads((WORK_DIR / "failure_chart_data.json").read_text(encoding="utf-8"))

MODELS_3 = ["GR00T Frozen LLM", "Pi0.5", "MolmoAct2"]
MODEL_LABEL = {
    "VLA-0": "VLA-0",
    "GR00T Full FT": "GR00T-N1.7 Full FT",
    "GR00T Frozen LLM": "GR00T Frozen LLM",
    "Pi0.5": "Pi0.5",
    "MolmoAct2": "MolmoAct2",
}


def quote_sheet(name: str) -> str:
    return "'" + name.replace("'", "''") + "'"


def get_row(rows: list[dict], **keys: object) -> dict:
    matches = [row for row in rows if all(row.get(key) == value for key, value in keys.items())]
    if len(matches) != 1:
        raise ValueError(f"Expected one row for {keys}; found {len(matches)}")
    return matches[0]


def set_row(ws, row: int, start_col: int, values: list[object]) -> None:
    for offset, value in enumerate(values):
        ws.cell(row=row, column=start_col + offset, value=value)


def custom_error_bars(sheet_name: str, minus_range: str, plus_range: str, direction: str = "y") -> ErrorBars:
    def source(address: str) -> NumDataSource:
        min_col, min_row, max_col, max_row = range_boundaries(address.replace("$", ""))
        values = [
            wb[sheet_name].cell(row=row, column=column).value
            for row in range(min_row, max_row + 1)
            for column in range(min_col, max_col + 1)
        ]
        cache = NumData(
            formatCode="0.0%",
            ptCount=len(values),
            pt=[NumVal(idx=index, v=float(value)) for index, value in enumerate(values) if value is not None],
        )
        return NumDataSource(numRef=NumRef(f=f"{quote_sheet(sheet_name)}!{address}", numCache=cache))

    error_bars = ErrorBars(
        errDir=direction,
        errBarType="both",
        errValType="cust",
        noEndCap=False,
        minus=source(minus_range),
        plus=source(plus_range),
    )
    error_bars.spPr = GraphicalProperties()
    error_bars.spPr.line.solidFill = "000000"
    error_bars.spPr.line.width = 12700
    return error_bars


def attach_three_errors(chart, sheet_name: str, row_start: int, row_end: int, error_cols: list[tuple[str, str]]) -> None:
    for series, (minus_col, plus_col) in zip(chart.ser, error_cols):
        series.errBars = custom_error_bars(
            sheet_name,
            f"${minus_col}${row_start}:${minus_col}${row_end}",
            f"${plus_col}${row_start}:${plus_col}${row_end}",
        )


def refresh_num_cache(series, values: list[float | None], format_code: str = "0.0%") -> None:
    points = [NumVal(idx=index, v=float(value)) for index, value in enumerate(values) if value is not None]
    series.val.numRef.numCache = NumData(formatCode=format_code, ptCount=len(values), pt=points)


def refresh_simple_labels(chart, categories: list[str], series_names: list[str]) -> None:
    for series, series_name in zip(chart.ser, series_names):
        if series.cat and series.cat.strRef:
            series.cat.strRef.strCache = StrData(
                ptCount=len(categories),
                pt=[StrVal(idx=index, v=value) for index, value in enumerate(categories)],
            )
        if series.tx and series.tx.strRef:
            series.tx.strRef.strCache = StrData(ptCount=1, pt=[StrVal(idx=0, v=series_name)])


def simple_chart_signature(chart) -> dict:
    anchor = chart.anchor
    anchor_signature = {
        "from": (anchor._from.col, anchor._from.colOff, anchor._from.row, anchor._from.rowOff),
    }
    if hasattr(anchor, "to"):
        anchor_signature["to"] = (anchor.to.col, anchor.to.colOff, anchor.to.row, anchor.to.rowOff)
    elif hasattr(anchor, "ext"):
        anchor_signature["ext"] = (anchor.ext.cx, anchor.ext.cy)
    series = []
    for item in chart.ser:
        value_formula = item.val.numRef.f if item.val and item.val.numRef else None
        category_formula = None
        if item.cat:
            if item.cat.strRef:
                category_formula = item.cat.strRef.f
            elif item.cat.multiLvlStrRef:
                category_formula = item.cat.multiLvlStrRef.f
            elif item.cat.numRef:
                category_formula = item.cat.numRef.f
        series.append((value_formula, category_formula))
    return {
        "class": type(chart).__name__,
        "type": getattr(chart, "type", None),
        "grouping": getattr(chart, "grouping", None),
        "width": chart.width,
        "height": chart.height,
        "anchor": anchor_signature,
        "series": series,
    }


def worksheet_signature(ws) -> dict:
    cells = []
    for (row, column), cell in sorted(ws._cells.items()):
        cells.append((row, column, cell.value, cell.data_type, cell.style_id, cell.number_format))
    return {
        "cells": cells,
        "charts": [simple_chart_signature(chart) for chart in ws._charts],
        "sheet_format": (ws.sheet_format.defaultColWidth, ws.sheet_format.defaultRowHeight),
        "freeze": str(ws.freeze_panes),
        "gridlines": ws.sheet_view.showGridLines,
    }


wb = load_workbook(SOURCE)
protected_sheets = [
    "Preliminary Task Results",
    "Pi0.5 Robustness",
    "Paraphrasing Comparison v2",
    " Paraphrasing Comparison v1",
]
protected_before = {name: worksheet_signature(wb[name]) for name in protected_sheets}


# Familiar spatial grounding.
ws = wb["Familiar Spatial Task Chart"]
rows = EXPECTED["Familiar Spatial Task Chart"]
tasks = [
    ("Counting", "Counting"),
    ("Ordinal Position", "Ordinal Reference Ordering"),
    ("Relational Placement", "Relational Placement"),
    ("Physical State", "State Recognition"),
    ("Relative Size", "Size Recognition"),
    ("Referential Description", "Referential Disambiguation"),
]
models = ["VLA-0", "GR00T Full FT", "GR00T Frozen LLM", "Pi0.5", "MolmoAct2"]
for col, model in enumerate(models, start=2):
    ws.cell(1, col, MODEL_LABEL[model])
for excel_row, (display_task, source_task) in enumerate(tasks, start=2):
    ws.cell(excel_row, 1, display_task)
    for index, model in enumerate(models):
        datum = get_row(rows, Task=source_task, Model=model)
        rate_col = 2 + index
        ci_col = 7 + 2 * index
        scored_col = 17 + index
        error_col = 22 + 2 * index
        set_row(ws, excel_row, rate_col, [datum["Rate"]])
        set_row(ws, excel_row, ci_col, [datum["CI Lower"], datum["CI Upper"]])
        set_row(ws, excel_row, scored_col, [datum["Scored"]])
        set_row(ws, excel_row, error_col, [datum["Lower Error"], datum["Upper Error"]])
chart = ws._charts[0]
if len(chart.ser) == 6:
    del chart.ser[5]
error_cols = [("V", "W"), ("X", "Y"), ("Z", "AA"), ("AB", "AC"), ("AD", "AE")]
for series, (minus_col, plus_col) in zip(chart.ser, error_cols):
    series.errBars = custom_error_bars("Familiar Spatial Task Chart", f"${minus_col}$2:${minus_col}$7", f"${plus_col}$2:${plus_col}$7")
for index, series in enumerate(chart.ser):
    refresh_num_cache(series, [ws.cell(row, 2 + index).value for row in range(2, 8)])
refresh_simple_labels(chart, [task[0] for task in tasks], [MODEL_LABEL[model] for model in models])


# Matched manipulation and novel spatial grounding task charts.
for sheet_name in ["Matched Manipulation Task Chart", "Novel Spatial Task Chart"]:
    ws = wb[sheet_name]
    rows = EXPECTED[sheet_name]
    set_row(ws, 1, 2, ["GR00T Frozen LLM", "Pi0.5", "MolmoAct2"])
    set_row(ws, 1, 14, [
        "GR00T Lower error", "GR00T Upper error", "Pi0.5 Lower error", "Pi0.5 Upper error", "MolmoAct2 Lower error", "MolmoAct2 Upper error"
    ])
    for excel_row, (display_task, source_task) in enumerate(tasks, start=2):
        ws.cell(excel_row, 1, display_task)
        data = [get_row(rows, Task=source_task, Model=model) for model in MODELS_3]
        set_row(ws, excel_row, 2, [datum["Rate"] for datum in data])
        set_row(ws, excel_row, 5, [value for datum in data for value in (datum["CI Lower"], datum["CI Upper"])])
        set_row(ws, excel_row, 11, [datum["Scored"] for datum in data])
        set_row(ws, excel_row, 14, [value for datum in data for value in (datum["Lower Error"], datum["Upper Error"])])
    chart = ws._charts[0]
    attach_three_errors(chart, sheet_name, 2, 7, [("N", "O"), ("P", "Q"), ("R", "S")])
    for index, series in enumerate(chart.ser):
        refresh_num_cache(series, [ws.cell(row, 2 + index).value for row in range(2, 8)])
    refresh_simple_labels(chart, [task[0] for task in tasks], ["GR00T Frozen LLM", "Pi0.5", "MolmoAct2"])


# Four no-movement diagnostics plus aggregate.
ws = wb["No Movement"]
rows = EXPECTED["No Movement"]
no_move_rows = [
    ("Counting: goal already satisfied", {"Task": "Counting", "Condition": "Goal already satisfied / no movement"}),
    ("Ordinal Position: invalid ordinal", {"Task": "Ordinal Reference Ordering", "Condition": "Invalid ordinal / no-movement diagnostic"}),
    ("Physical State: no valid target", {"Task": "State Recognition", "Condition": "No valid target / no movement"}),
    ("Referential Description: no valid candidate", {"Task": "Referential Disambiguation", "Condition": "No valid candidate category / no movement"}),
    ("Aggregate across four tasks", {"Task": "Aggregate across four tasks", "Condition": "Aggregate"}),
]
for excel_row, (label, keys) in enumerate(no_move_rows, start=2):
    ws.cell(excel_row, 1, label)
    data = [get_row(rows, Model=model, **keys) for model in MODELS_3]
    set_row(ws, excel_row, 2, [datum["Rate"] for datum in data])
    set_row(ws, excel_row, 5, [value for datum in data for value in (datum["CI Lower"], datum["CI Upper"])])
    set_row(ws, excel_row, 11, [datum["Scored"] for datum in data])
    set_row(ws, excel_row, 14, [value for datum in data for value in (datum["Lower Error"], datum["Upper Error"])])
chart = ws._charts[0]
attach_three_errors(chart, "No Movement", 2, 6, [("N", "O"), ("P", "Q"), ("R", "S")])
for index, series in enumerate(chart.ser):
    refresh_num_cache(series, [ws.cell(row, 2 + index).value for row in range(2, 7)])
refresh_simple_labels(chart, [row[0] for row in no_move_rows], [ws.cell(1, col).value for col in range(2, 5)])


# Relational no-movement versus corrective movement.
ws = wb["Relational Movement No Movement"]
rows = EXPECTED["Relational Movement No Movement"]
set_row(ws, 1, 14, [
    "GR00T Lower error", "GR00T Upper error", "Pi0.5 Lower error", "Pi0.5 Upper error", "MolmoAct2 Lower error", "MolmoAct2 Upper error"
])
relational_rows = [
    ("Relational goal already satisfied / no movement", "Goal already satisfied / no movement"),
    ("Relational corrective movement required", "Violated-relation correction control"),
]
for excel_row, (label, condition) in enumerate(relational_rows, start=2):
    ws.cell(excel_row, 1, label)
    data = [get_row(rows, Condition=condition, Model=model) for model in MODELS_3]
    set_row(ws, excel_row, 2, [datum["Rate"] for datum in data])
    set_row(ws, excel_row, 5, [value for datum in data for value in (datum["CI Lower"], datum["CI Upper"])])
    set_row(ws, excel_row, 11, [datum["Scored"] for datum in data])
    set_row(ws, excel_row, 14, [value for datum in data for value in (datum["Lower Error"], datum["Upper Error"])])
chart = ws._charts[0]
attach_three_errors(chart, "Relational Movement No Movement", 2, 3, [("N", "O"), ("P", "Q"), ("R", "S")])
for index, series in enumerate(chart.ser):
    refresh_num_cache(series, [ws.cell(row, 2 + index).value for row in range(2, 4)])
refresh_simple_labels(chart, [row[0] for row in relational_rows], [ws.cell(1, col).value for col in range(2, 5)])


# Standard versus cross-task objects.
ws = wb["Cross Task Object Chart v2"]
rows = EXPECTED["Cross Task Object Chart v2"]
set_row(ws, 1, 14, [
    "GR00T Lower error", "GR00T Upper error", "Pi0.5 Lower error", "Pi0.5 Upper error", "MolmoAct2 Lower error", "MolmoAct2 Upper error"
])
categories = [
    "Familiar spatial grounding, standard objects",
    "Familiar spatial grounding, cross-task objects",
]
for excel_row, category in enumerate(categories, start=2):
    data = [get_row(rows, Category=category, Model=model) for model in MODELS_3]
    ws.cell(excel_row, 1, category)
    set_row(ws, excel_row, 2, [datum["Rate"] for datum in data])
    set_row(ws, excel_row, 5, [value for datum in data for value in (datum["CI Lower"], datum["CI Upper"])])
    set_row(ws, excel_row, 11, [datum["Scored"] for datum in data])
    set_row(ws, excel_row, 14, [value for datum in data for value in (datum["Lower Error"], datum["Upper Error"])])
chart = ws._charts[0]
attach_three_errors(chart, "Cross Task Object Chart v2", 2, 3, [("N", "O"), ("P", "Q"), ("R", "S")])
for index, series in enumerate(chart.ser):
    refresh_num_cache(series, [ws.cell(row, 2 + index).value for row in range(2, 4)])
refresh_simple_labels(chart, categories, [ws.cell(1, col).value for col in range(2, 5)])


# Original versus paraphrased instructions on the same 60 scenes.
ws = wb["Paraphrasing Comparison v3"]
rows = EXPECTED["Paraphrasing Comparison v3"]
set_row(ws, 1, 2, ["GR00T Frozen LLM", "Pi0.5", "MolmoAct2"])
main_rows = [
    (2, "Matched Manipulation", "Original"),
    (3, "Matched Manipulation", "Paraphrased"),
    (4, "Novel Spatial Grounding", "Original"),
    (5, "Novel Spatial Grounding", "Paraphrased"),
]
for excel_row, group, phrasing in main_rows:
    data = [get_row(rows, Group=group, Model=model, Phrasing=phrasing) for model in MODELS_3]
    set_row(ws, excel_row, 2, [datum["Rate"] for datum in data])
    set_row(ws, excel_row, 5, [value for datum in data for value in (datum["CI Lower"], datum["CI Upper"])])
    set_row(ws, excel_row, 11, [datum["Successes"] for datum in data])
    set_row(ws, excel_row, 14, [datum["Scored"] for datum in data])
headers = ["Group", "Model", "Original", "Paraphrased", "Original Lower error", "Original Upper error", "Paraphrased Lower error", "Paraphrased Upper error"]
set_row(ws, 1, 17, headers)
helper_rows: list[list[object]] = []
for group in ["Matched Manipulation", "Novel Spatial Grounding"]:
    for model in MODELS_3:
        original = get_row(rows, Group=group, Model=model, Phrasing="Original")
        paraphrased = get_row(rows, Group=group, Model=model, Phrasing="Paraphrased")
        helper_rows.append([
            group, MODEL_LABEL[model], original["Rate"], paraphrased["Rate"],
            original["Lower Error"], original["Upper Error"], paraphrased["Lower Error"], paraphrased["Upper Error"],
        ])
    if group == "Matched Manipulation":
        helper_rows.append([None] * 8)
for row_number, values in enumerate(helper_rows, start=2):
    set_row(ws, row_number, 17, values)
chart = ws._charts[0]
chart.ser[0].errBars = custom_error_bars("Paraphrasing Comparison v3", "$U$2:$U$8", "$V$2:$V$8")
chart.ser[1].errBars = custom_error_bars("Paraphrasing Comparison v3", "$W$2:$W$8", "$X$2:$X$8")
refresh_num_cache(chart.ser[0], [ws.cell(row, 19).value for row in range(2, 9)])
refresh_num_cache(chart.ser[1], [ws.cell(row, 20).value for row in range(2, 9)])
for series in chart.ser:
    if series.cat and series.cat.multiLvlStrRef:
        series.cat.multiLvlStrRef.multiLvlStrCache = None


# Overall comparison chart.
ws = wb["Emergent Skills Data"]
rows = EXPECTED["Emergent Skills Data"]
set_row(ws, 1, 2, ["GR00T Frozen LLM", "Pi0.5", "MolmoAct2"])
set_row(ws, 1, 5, [
    "GR00T Successes", "Pi0.5 Successes", "MolmoAct2 Successes",
    "GR00T Scored", "Pi0.5 Scored", "MolmoAct2 Scored",
    "GR00T CI lower", "GR00T CI upper", "Pi0.5 CI lower", "Pi0.5 CI upper", "MolmoAct2 CI lower", "MolmoAct2 CI upper",
    "GR00T Lower error", "GR00T Upper error", "Pi0.5 Lower error", "Pi0.5 Upper error", "MolmoAct2 Lower error", "MolmoAct2 Upper error",
])
tracks = [
    ("Matched Manipulation Control", "Matched Manipulation Control"),
    ("Familiar Spatial Grounding", "Familiar Spatial Concepts"),
    ("Novel Spatial Grounding", "Novel Spatial Concepts"),
]
for excel_row, (label, track) in enumerate(tracks, start=2):
    data = [get_row(rows, **{"Evaluation track": track, "Model": model}) for model in MODELS_3]
    set_row(ws, excel_row, 1, [label, *[datum["Rate"] for datum in data]])
    set_row(ws, excel_row, 5, [datum["Successes"] for datum in data])
    set_row(ws, excel_row, 8, [datum["Scored"] for datum in data])
    set_row(ws, excel_row, 11, [value for datum in data for value in (datum["CI Lower"], datum["CI Upper"])])
    set_row(ws, excel_row, 17, [value for datum in data for value in (datum["Lower Error"], datum["Upper Error"])])
chart = ws._charts[0]
attach_three_errors(chart, "Emergent Skills Data", 2, 4, [("Q", "R"), ("S", "T"), ("U", "V")])
for index, series in enumerate(chart.ser):
    refresh_num_cache(series, [ws.cell(row, 2 + index).value for row in range(2, 5)])
refresh_simple_labels(chart, [track[0] for track in tracks], ["GR00T Frozen LLM", "Pi0.5", "MolmoAct2"])


# Detailed overall results table and localize its stale external chart references.
ws = wb["Sheet2"]
rows = EXPECTED["Sheet2"]
row_map = [
    (5, "Novel Spatial Concepts", "MolmoAct2"), (6, "Novel Spatial Concepts", "Pi0.5"), (7, "Novel Spatial Concepts", "GR00T Frozen LLM"),
    (8, "Familiar Spatial Concepts", "MolmoAct2"), (9, "Familiar Spatial Concepts", "Pi0.5"), (10, "Familiar Spatial Concepts", "GR00T Frozen LLM"),
    (11, "Matched Manipulation Control", "MolmoAct2"), (12, "Matched Manipulation Control", "Pi0.5"), (13, "Matched Manipulation Control", "GR00T Frozen LLM"),
]
for excel_row, track, model in row_map:
    datum = get_row(rows, **{"Evaluation track": track, "Model": model})
    display_model = "π0.5" if model == "Pi0.5" else MODEL_LABEL[model]
    set_row(ws, excel_row, 2, [
        display_model, "Complete benchmark", datum["Successes"], datum["Scored"] - datum["Successes"], datum["Scored"], datum["Rate"],
        datum["CI Lower"], datum["CI Upper"], datum["Lower Error"], datum["Upper Error"],
        f"{datum['Rate']:.1%} [{datum['CI Lower']:.1%}–{datum['CI Upper']:.1%}]",
    ])
helper_headers = ["Evaluation track", "MolmoAct2", "π0.5", "GR00T Frozen LLM", "Molmo lower", "Molmo upper", "Pi lower", "Pi upper", "GR00T lower", "GR00T upper"]
set_row(ws, 4, 17, helper_headers)
for excel_row, track in enumerate(["Novel Spatial Concepts", "Familiar Spatial Concepts", "Matched Manipulation Control"], start=5):
    data = [get_row(rows, **{"Evaluation track": track, "Model": model}) for model in ["MolmoAct2", "Pi0.5", "GR00T Frozen LLM"]]
    set_row(ws, excel_row, 17, [track, *[datum["Rate"] for datum in data], *[value for datum in data for value in (datum["Lower Error"], datum["Upper Error"])]] )
chart = ws._charts[0]
value_cols = ["R", "S", "T"]
error_cols = [("U", "V"), ("W", "X"), ("Y", "Z")]
for index, series in enumerate(chart.ser):
    series.val.numRef.f = f"'Sheet2'!${value_cols[index]}$5:${value_cols[index]}$7"
    series.cat.strRef.f = "'Sheet2'!$Q$5:$Q$7"
    series.errBars = custom_error_bars("Sheet2", f"${error_cols[index][0]}$5:${error_cols[index][0]}$7", f"${error_cols[index][1]}$5:${error_cols[index][1]}$7")
    refresh_num_cache(series, [ws.cell(row, 18 + index).value for row in range(5, 8)], format_code="0.0%")
chart.ser[0].tx.v = "MolmoAct2"
chart.ser[1].tx.v = "π0.5"
chart.ser[2].tx.v = "GR00T Frozen LLM"
refresh_simple_labels(chart, ["Novel Spatial Concepts", "Familiar Spatial Concepts", "Matched Manipulation Control"], ["MolmoAct2", "π0.5", "GR00T Frozen LLM"])


def add_failure_sheet(sheet_name: str, title: str, payload: dict) -> None:
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False
    rows = sorted(payload["rows"], key=lambda row: row["Count"], reverse=True)
    ws["A1"] = title
    ws["A2"] = f"{payload['scope']}. Shares and Wilson intervals are conditional on failure (n={payload['total']})."
    ws["A3"] = "Source: Spa_Bench_All_Failure_Audit.xlsx."
    headers = ["Failure mode", "Count", "Share of failures", "Wilson 95% lower", "Wilson 95% upper", "Lower error", "Upper error"]
    set_row(ws, 5, 1, headers)
    for excel_row, row in enumerate(rows, start=6):
        set_row(ws, excel_row, 1, [
            row["Failure Mode"], row["Count"], row["Share"], row["CI Lower"], row["CI Upper"],
            row["Share"] - row["CI Lower"], row["CI Upper"] - row["Share"],
        ])
    ws["A1"].font = Font(name="Arial", size=14, bold=True, color="1F1F1F")
    for cell in ws[2] + ws[3]:
        cell.font = Font(name="Arial", size=10, italic=True, color="595959")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[5]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="D9E2F3")
    for row_cells in ws.iter_rows(min_row=6, max_row=5 + len(rows), min_col=1, max_col=7):
        for cell in row_cells:
            cell.font = Font(name="Arial", size=10, color="1F1F1F")
            cell.border = Border(bottom=thin)
    for row_number in range(6, 6 + len(rows)):
        ws.cell(row_number, 2).number_format = "#,##0"
        for col in range(3, 8):
            ws.cell(row_number, col).number_format = "0.0%"
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 11
    for col in ["C", "D", "E", "F", "G"]:
        ws.column_dimensions[col].width = 17
    chart = BarChart()
    chart.type = "bar"
    chart.grouping = "clustered"
    chart.style = 10
    chart.title = title
    chart.legend = None
    chart.width = 15
    chart.height = 7.5
    chart.x_axis.scaling.min = 0
    chart.x_axis.scaling.max = 0.6
    chart.x_axis.numFmt = "0%"
    chart.x_axis.majorGridlines = copy(chart.x_axis.majorGridlines)
    data = Reference(ws, min_col=3, min_row=5, max_row=5 + len(rows))
    categories = Reference(ws, min_col=1, min_row=6, max_row=5 + len(rows))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.ser[0].graphicalProperties.solidFill = "4F81BD"
    chart.ser[0].graphicalProperties.line.solidFill = "4F81BD"
    chart.ser[0].errBars = custom_error_bars(sheet_name, f"$F$6:$F${5 + len(rows)}", f"$G$6:$G${5 + len(rows)}", direction="x")
    ws.add_chart(chart, "A16")
    ws.print_area = "A1:T40"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1


add_failure_sheet("Failure Analysis - Novel", "Failure modes in novel spatial grounding", FAILURES["Novel"])
add_failure_sheet("Failure Analysis - All", "Failure modes across all failed rollouts", FAILURES["All"])


OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"
wb.save(OUTPUT)


# Re-open and verify that protected sheets are unchanged and requested output exists.
verified = load_workbook(OUTPUT, data_only=False)
protected_after = {name: worksheet_signature(verified[name]) for name in protected_sheets}
protected_equal = {name: protected_before[name] == protected_after[name] for name in protected_sheets}
if not all(protected_equal.values()):
    changed = [name for name, equal in protected_equal.items() if not equal]
    raise AssertionError(f"Protected sheets changed: {changed}")

verification = {
    "output": str(OUTPUT),
    "protected_sheets_unchanged": protected_equal,
    "chart_count": sum(len(ws._charts) for ws in verified.worksheets),
    "failure_sheets": {
        name: {"rows": verified[name].max_row, "charts": len(verified[name]._charts)}
        for name in ["Failure Analysis - Novel", "Failure Analysis - All"]
    },
    "corrected_cells": {
        "Familiar Spatial Task Chart!C2": verified["Familiar Spatial Task Chart"]["C2"].value,
        "Familiar Spatial Task Chart!R2": verified["Familiar Spatial Task Chart"]["R2"].value,
        "Familiar Spatial Task Chart!F6": verified["Familiar Spatial Task Chart"]["F6"].value,
        "Cross Task Object Chart v2!D2": verified["Cross Task Object Chart v2"]["D2"].value,
        "Paraphrasing Comparison v3!C3": verified["Paraphrasing Comparison v3"]["C3"].value,
        "Sheet2!D8": verified["Sheet2"]["D8"].value,
    },
}
(WORK_DIR / "fallback_verification.json").write_text(json.dumps(verification, indent=2), encoding="utf-8")
print(json.dumps(verification, indent=2))
