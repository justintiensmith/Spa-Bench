from __future__ import annotations

import json
from collections import defaultdict
from html import escape
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


WORKBOOK = Path(
    "/Users/justintiensmith/Documents/Consolidated_Model_Evaluations_with_Custom_Error_Bars.xlsx"
)
OUTPUT_DIR = Path(
    "/Users/justintiensmith/Documents/lerobot/output/pdf/"
    "ICRA_paper_v3_main_appendix_2026-09-06/Figures"
)

MODEL_SHEETS = {
    "GR00T Frozen LLM": "GR00T Vision",
    r"$\pi_{0.5}$": "Pi0.5",
    "MolmoAct2": "MolmoAct2",
}

MODEL_COLORS = {
    "GR00T Frozen LLM": "#C95A48",
    r"$\pi_{0.5}$": "#6CA363",
    "MolmoAct2": "#5E82E6",
}

TASKS = [
    "Counting",
    "Ordinal Reference Ordering",
    "Relational Placement",
    "State Recognition",
    "Size Recognition",
    "Referential Disambiguation",
]

NOVEL = "Concept/compositional transfer"
DIRECT = "Matched manipulation control"
PARAPHRASE = "Language robustness"
SUCCESS = {"Success": 1.0, "Failure": 0.0}
B = 100_000
SEED = 20260906


def load_rows(sheet_name: str) -> list[dict[str, object]]:
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    sheet = workbook[sheet_name]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))][:40]
    rows = []
    for values in sheet.iter_rows(min_row=2, max_col=40, values_only=True):
        row = dict(zip(headers, values))
        if row["Status"] in SUCCESS and row["Task"] in TASKS:
            rows.append(row)
    workbook.close()
    return rows


def validate_reference(
    row: dict[str, object],
    reference: dict[str, object],
    expected_condition: str,
) -> None:
    assert reference["Condition"] == expected_condition
    assert row["Task"] == reference["Task"]
    assert row["Scene Family"] == reference["Scene Family"]


def construct_pairs(rows: list[dict[str, object]]) -> dict[str, dict[str, np.ndarray]]:
    indexed = {(row["Task"], int(row["Episode"])): row for row in rows}
    effects: dict[str, dict[str, list[float]]] = {
        "direct_minus_novel": defaultdict(list),
        "novel_paraphrase_minus_original": defaultdict(list),
        "direct_paraphrase_minus_original": defaultdict(list),
    }

    for task in TASKS:
        task_rows = [row for row in rows if row["Task"] == task]
        direct_rows = [row for row in task_rows if row["Condition"] == DIRECT]
        paraphrase_rows = [row for row in task_rows if row["Condition"] == PARAPHRASE]
        assert len(direct_rows) == 20, (task, len(direct_rows))
        assert len(paraphrase_rows) == 20, (task, len(paraphrase_rows))

        for direct_row in direct_rows:
            key = (task, int(direct_row["Reference Episode"]))
            novel_row = indexed[key]
            validate_reference(direct_row, novel_row, NOVEL)
            effects["direct_minus_novel"][task].append(
                SUCCESS[str(direct_row["Status"])] - SUCCESS[str(novel_row["Status"])]
            )

        for paraphrase_row in paraphrase_rows:
            key = (task, int(paraphrase_row["Reference Episode"]))
            original_row = indexed[key]
            assert original_row["Condition"] in (NOVEL, DIRECT)
            validate_reference(paraphrase_row, original_row, str(original_row["Condition"]))
            effect = SUCCESS[str(paraphrase_row["Status"])] - SUCCESS[str(original_row["Status"])]
            target = (
                "novel_paraphrase_minus_original"
                if original_row["Condition"] == NOVEL
                else "direct_paraphrase_minus_original"
            )
            effects[target][task].append(effect)

        assert len(effects["direct_minus_novel"][task]) == 20
        assert len(effects["novel_paraphrase_minus_original"][task]) == 10
        assert len(effects["direct_paraphrase_minus_original"][task]) == 10

    return {
        comparison: {task: np.asarray(values, dtype=float) for task, values in by_task.items()}
        for comparison, by_task in effects.items()
    }


def paired_bootstrap(
    by_task: dict[str, np.ndarray],
    rng: np.random.Generator,
) -> tuple[float, float, float]:
    point = float(np.mean([np.mean(by_task[task]) for task in TASKS]))
    replicate_task_means = []
    for task in TASKS:
        values = by_task[task]
        indices = rng.integers(0, len(values), size=(B, len(values)))
        replicate_task_means.append(values[indices].mean(axis=1))
    replicates = np.vstack(replicate_task_means).mean(axis=0)
    lower, upper = np.quantile(replicates, [0.025, 0.975])
    return point, float(lower), float(upper)


def compute_results() -> dict[str, dict[str, tuple[float, float, float]]]:
    rng = np.random.default_rng(SEED)
    results: dict[str, dict[str, tuple[float, float, float]]] = {}
    for model, sheet_name in MODEL_SHEETS.items():
        pairs = construct_pairs(load_rows(sheet_name))
        results[model] = {
            comparison: paired_bootstrap(by_task, rng)
            for comparison, by_task in pairs.items()
        }
    return results


def format_signed(value: float) -> str:
    if abs(value) < 0.0005:
        return "0.0"
    return f"{value * 100:+.1f}".replace("-", "−")


def draw_figure(results: dict[str, dict[str, tuple[float, float, float]]]) -> None:
    width, height = 2200, 740
    plot_top, plot_bottom = 160, 585
    panel_specs = [(385, 730), (1370, 760)]
    x_min, x_max = -40.0, 60.0
    ticks = [-30, -15, 0, 15, 30, 45, 60]
    models = list(MODEL_SHEETS)
    y_positions = [255, 385, 515]

    def sx(value: float, x0: float, panel_width: float) -> float:
        return x0 + (value - x_min) / (x_max - x_min) * panel_width

    def line(x1: float, y1: float, x2: float, y2: float, **attrs: object) -> str:
        properties = " ".join(f'{key.replace("_", "-")}="{value}"' for key, value in attrs.items())
        return f'<line x1="{x1:.2f}" y1="{y1:.2f}" x2="{x2:.2f}" y2="{y2:.2f}" {properties}/>'

    def text(x: float, y: float, value: str, **attrs: object) -> str:
        properties = " ".join(f'{key.replace("_", "-")}="{value_}"' for key, value_ in attrs.items())
        return f'<text x="{x:.2f}" y="{y:.2f}" {properties}>{escape(value)}</text>'

    def whisker(
        x0: float,
        panel_width: float,
        y: float,
        point: float,
        lower: float,
        upper: float,
        color: str,
        marker: str,
    ) -> list[str]:
        xp, xl, xu = (sx(value, x0, panel_width) for value in (point, lower, upper))
        elements = [
            line(xl, y, xu, y, stroke=color, stroke_width=8, stroke_linecap="round"),
            line(xl, y - 12, xl, y + 12, stroke=color, stroke_width=5),
            line(xu, y - 12, xu, y + 12, stroke=color, stroke_width=5),
        ]
        if marker == "circle":
            elements.append(
                f'<circle cx="{xp:.2f}" cy="{y:.2f}" r="11" fill="{color}" stroke="#FFFFFF" stroke-width="3"/>'
            )
        else:
            elements.append(
                f'<rect x="{xp - 10:.2f}" y="{y - 10:.2f}" width="20" height="20" fill="{color}" stroke="#FFFFFF" stroke-width="3"/>'
            )
        return elements

    svg: list[str] = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="7.15in" height="2.40in" viewBox="0 0 {width} {height}">',
        '<rect width="100%" height="100%" fill="#FFFFFF"/>',
        '<g font-family="Arial, Helvetica, sans-serif" fill="#111111">',
    ]

    titles = ["(a) Direct manipulation advantage", "(b) Effect of paraphrasing"]
    for (x0, panel_width), title in zip(panel_specs, titles):
        svg.append(text(x0, 57, title, font_size=39, font_weight="700"))
        for tick in ticks:
            x_tick = sx(tick, x0, panel_width)
            is_zero = tick == 0
            svg.append(
                line(
                    x_tick,
                    plot_top,
                    x_tick,
                    plot_bottom,
                    stroke="#555555" if is_zero else "#D9D9D9",
                    stroke_width=4 if is_zero else 2,
                )
            )
            svg.append(
                text(
                    x_tick,
                    628,
                    str(tick).replace("-", "−"),
                    font_size=28,
                    text_anchor="middle",
                )
            )
        svg.append(line(x0, plot_bottom, x0 + panel_width, plot_bottom, stroke="#777777", stroke_width=2))
        svg.append(
            text(
                x0 + panel_width / 2,
                690,
                "Difference in success rate (percentage points)",
                font_size=29,
                text_anchor="middle",
                font_weight="600",
            )
        )

    x0, panel_width = panel_specs[0]
    for y, model in zip(y_positions, models):
        color = MODEL_COLORS[model]
        label = "π₀.₅" if model == r"$\pi_{0.5}$" else model
        svg.append(
            text(
                x0 - 28,
                y + 10,
                label,
                font_size=29,
                text_anchor="end",
                font_weight="700",
                fill=color,
            )
        )
        point, lower, upper = results[model]["direct_minus_novel"]
        point, lower, upper = 100 * point, 100 * lower, 100 * upper
        svg.extend(whisker(x0, panel_width, y, point, lower, upper, color, "circle"))
        svg.append(
            text(
                sx(upper, x0, panel_width) + 16,
                y + 9,
                format_signed(point / 100),
                font_size=27,
                font_weight="600",
            )
        )

    x0, panel_width = panel_specs[1]
    legend_y = 110
    svg.append(f'<circle cx="{x0 + 10}" cy="{legend_y}" r="10" fill="#666666"/>')
    svg.append(text(x0 + 31, legend_y + 9, "Novel spatial", font_size=26))
    svg.append(f'<rect x="{x0 + 255}" y="{legend_y - 10}" width="20" height="20" fill="#666666"/>')
    svg.append(text(x0 + 288, legend_y + 9, "Direct manipulation", font_size=26))

    comparison_keys = {
        "novel": "novel_paraphrase_minus_original",
        "direct": "direct_paraphrase_minus_original",
    }
    offsets = {"novel": -18, "direct": 18}
    markers = {"novel": "circle", "direct": "square"}
    for y, model in zip(y_positions, models):
        color = MODEL_COLORS[model]
        label = "π₀.₅" if model == r"$\pi_{0.5}$" else model
        svg.append(
            text(
                x0 - 28,
                y + 10,
                label,
                font_size=29,
                text_anchor="end",
                font_weight="700",
                fill=color,
            )
        )
        for instruction_type in ("novel", "direct"):
            point, lower, upper = results[model][comparison_keys[instruction_type]]
            point, lower, upper = 100 * point, 100 * lower, 100 * upper
            y_mark = y + offsets[instruction_type]
            svg.extend(
                whisker(
                    x0,
                    panel_width,
                    y_mark,
                    point,
                    lower,
                    upper,
                    color,
                    markers[instruction_type],
                )
            )
            svg.append(
                text(
                    sx(upper, x0, panel_width) + 14,
                    y_mark + 8,
                    format_signed(point / 100),
                    font_size=24,
                    font_weight="600",
                )
            )

    svg.extend(["</g>", "</svg>"])
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    (OUTPUT_DIR / "same_scene_effects.svg").write_text("\n".join(svg), encoding="utf-8")


if __name__ == "__main__":
    computed = compute_results()
    serializable = {
        model: {
            comparison: {
                "estimate": estimate,
                "ci_lower": lower,
                "ci_upper": upper,
            }
            for comparison, (estimate, lower, upper) in comparisons.items()
        }
        for model, comparisons in computed.items()
    }
    print(json.dumps(serializable, indent=2))
    draw_figure(computed)
